#!/usr/bin/env python3
"""根据 Excel 配置批量调用 split_pdf_pages.py 进行 PDF 切割。

输入参数
--------
- excel_path: Excel 文件路径（.xlsx）。
- pdf_path: 源 PDF 文件路径。
- target_path: 绝对输出根目录。

Excel 规则（默认读取第一张表）：
- 第一列：relative_path（相对 target_path 的子路径）
- 第二列：start_page（起始页，1-based）
- 第三列：end_page（结束页，1-based，包含结束页）

每一行会执行一次：
python process/split_pdf_pages.py pdf_path start_page end_page target_path relative_path
"""

from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
	"""解析命令行参数。"""
	parser = argparse.ArgumentParser(
		description="Batch split a PDF by row rules from an Excel file."
	)
	parser.add_argument("excel_path", type=Path, help="Path to the Excel file (.xlsx).")
	parser.add_argument("pdf_path", type=Path, help="Path to the source PDF file.")
	parser.add_argument("target_path", type=Path, help="Absolute output root directory.")
	parser.add_argument(
		"--sheet",
		default=None,
		help="Sheet name to read. If omitted, the first worksheet is used.",
	)
	return parser.parse_args()


def parse_page_number(value: Any) -> int:
	"""将 Excel 单元格值解析为正整数页码。"""
	if value is None:
		raise ValueError("page is empty")

	if isinstance(value, int):
		page = value
	elif isinstance(value, float):
		if not value.is_integer():
			raise ValueError(f"page must be integer, got {value}")
		page = int(value)
	else:
		text = str(value).strip()
		if not text:
			raise ValueError("page is empty")
		page = int(text)

	if page < 1:
		raise ValueError(f"page must be >= 1, got {page}")
	return page


def main() -> int:
	"""程序入口。"""
	args = parse_args()

	excel_path = args.excel_path
	pdf_path = args.pdf_path
	target_path = args.target_path

	if not excel_path.exists() or not excel_path.is_file():
		print(f"Error: Excel file not found: {excel_path}", file=sys.stderr)
		return 1

	if excel_path.suffix.lower() != ".xlsx":
		print("Error: excel_path must be an .xlsx file.", file=sys.stderr)
		return 1

	if not pdf_path.exists() or not pdf_path.is_file():
		print(f"Error: PDF file not found: {pdf_path}", file=sys.stderr)
		return 1

	if pdf_path.suffix.lower() != ".pdf":
		print("Error: pdf_path must be a .pdf file.", file=sys.stderr)
		return 1

	if not target_path.is_absolute():
		print("Error: target_path must be an absolute path.", file=sys.stderr)
		return 1

	split_script = Path(__file__).resolve().parent / "process" / "split_pdf_pages.py"
	if not split_script.exists():
		print(f"Error: split script not found: {split_script}", file=sys.stderr)
		return 1

	try:
		workbook = load_workbook(excel_path, read_only=True, data_only=True)
	except Exception as exc:
		print(f"Error: failed to open Excel: {exc}", file=sys.stderr)
		return 1

	try:
		if args.sheet:
			if args.sheet not in workbook.sheetnames:
				print(
					f"Error: sheet '{args.sheet}' not found. Available: {workbook.sheetnames}",
					file=sys.stderr,
				)
				return 1
			worksheet = workbook[args.sheet]
		else:
			worksheet = workbook[workbook.sheetnames[0]]

		success_count = 0
		fail_count = 0

		for row_index, row in enumerate(
			worksheet.iter_rows(min_col=1, max_col=3, values_only=True),
			start=1,
		):
			relative_cell, start_cell, end_cell = row

			# 跳过完全空行
			if relative_cell is None and start_cell is None and end_cell is None:
				continue

			relative_text = "" if relative_cell is None else str(relative_cell).strip()
			if not relative_text:
				print(f"[Row {row_index}] Skip: relative_path is empty.")
				fail_count += 1
				continue

			try:
				start_page = parse_page_number(start_cell)
				end_page = parse_page_number(end_cell)
			except Exception as exc:
				print(f"[Row {row_index}] Skip: invalid page number: {exc}")
				fail_count += 1
				continue

			if start_page > end_page:
				print(
					f"[Row {row_index}] Skip: start_page {start_page} > end_page {end_page}."
				)
				fail_count += 1
				continue

			command = [
				sys.executable,
				str(split_script),
				str(pdf_path),
				str(start_page),
				str(end_page),
				str(target_path),
				relative_text,
			]
			result = subprocess.run(command, capture_output=True, text=True)

			if result.returncode == 0:
				success_count += 1
				message = result.stdout.strip() or "Success"
				print(f"[Row {row_index}] OK: {message}")
			else:
				fail_count += 1
				error_message = (result.stderr or result.stdout).strip() or "Unknown error"
				print(f"[Row {row_index}] FAIL: {error_message}")

		print("\nBatch finished.")
		print(f"Success: {success_count}")
		print(f"Failed: {fail_count}")

		return 0 if fail_count == 0 else 2
	finally:
		workbook.close()


if __name__ == "__main__":
	raise SystemExit(main())
