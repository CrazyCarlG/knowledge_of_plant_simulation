from openpyxl import load_workbook
import subprocess
import sys
from pathlib import Path

# -------------------------- 配置常量，集中修改 --------------------------
EXCEL_FILE = Path(r"C:\Users\z004bjuu\Downloads\all-objects-sections_allobjects.xlsx")
SCRIPT_PATH = Path(r"C:\Users\z004bjuu\Documents\knowledge_of_plant_simulation\scripts\pdf-to-knowledge\process\openclaude_readme_session.py")
BASE_DIR = Path(r"C:\Users\z004bjuu\Documents\knowledge_of_plant_simulation\01-plant-simulation-help")
DEBUG_BREAK_AFTER_FIRST_ROW = False   # True只跑第一行调试；False批量跑全部行
# ------------------------------------------------------------------------


def read_excel_rows(excel_path: Path):
    if not SCRIPT_PATH.exists():
        print(f"Error: 子脚本不存在: {SCRIPT_PATH}")
        return

    wb = None
    try:
        wb = load_workbook(filename=str(excel_path), read_only=True)
        ws = wb.active

        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            col_a = row[0]
            col_b = row[1]

            if row_idx == 1:
                print(f"第{row_idx}行：表头，跳过")
                continue

            # 跳过A列为空的无效行
            if col_a is None or str(col_a).strip() == "":
                print(f"第{row_idx}行：A列为空，跳过")
                continue
            col_a = str(col_a).strip().replace("\\", "/")
            file_path = BASE_DIR / col_a
            print(f"\n===== 第{row_idx}行 =====")
            print(f"A列拼接完整路径: {file_path}")
            print(f"B列内容: {col_b}")

            if not file_path.exists():
                print(f"第{row_idx}行：路径不存在，跳过")
                continue

            cmd = [
                sys.executable,
                str(SCRIPT_PATH),
                str(file_path),
            ]

            try:
                print(cmd)

                result = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    check=True
                )
                if result.stdout:
                    print("脚本stdout输出:\n", result.stdout)
                if result.stderr:
                    print("脚本stderr输出:\n", result.stderr)

                if DEBUG_BREAK_AFTER_FIRST_ROW:
                    print("\n【调试模式】仅处理第一行有效数据，终止循环")
                    break

            except subprocess.CalledProcessError as e:
                # 子脚本执行报错，打印stderr，不终止整个循环
                print(f"⚠️ 第{row_idx}行脚本执行失败！")
                print(f"Return code: {e.returncode}")
                print(f"STDERR:\n{e.stderr}")
                continue
            except Exception as e:
                print(f"⚠️ 第{row_idx}行发生未知异常: {str(e)}")
                continue

            # 调试模式：处理完第一行直接退出循环
            if DEBUG_BREAK_AFTER_FIRST_ROW:
                print("\n【调试模式】仅处理第一行，终止循环")
                break
    finally:
        # 无论正常/异常，一定关闭workbook释放文件
        if wb is not None:
            wb.close()

    print("\n✅ Excel全部行处理完成")


if __name__ == "__main__":
    read_excel_rows(EXCEL_FILE)